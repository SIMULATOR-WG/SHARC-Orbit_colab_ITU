"""Tests for the per-run result artifacts (R11–R15, R17, R22, R24)."""
from __future__ import annotations

import csv
import io
import json
import sys
from pathlib import Path

import numpy as np
import pytest

REPO = Path(__file__).resolve().parents[2]
for p in (str(REPO), str(REPO / "streamlit_app")):
    if p not in sys.path:
        sys.path.insert(0, p)

from src.epfd_stream_accumulator import EPFDStreamAccumulator  # noqa: E402
from src.epfd_calculator import check_article22_compliance  # noqa: E402
from streamlit_app.lib.result_artifacts import write_run_artifacts  # noqa: E402
from streamlit_app.lib.report import write_summary_html  # noqa: E402
from streamlit_app.lib.exports import run_to_xlsx  # noqa: E402


def _acc(n: int = 500) -> EPFDStreamAccumulator:
    acc = EPFDStreamAccumulator()
    rng = np.random.default_rng(7)
    for k in range(n):
        acc.add(time_s=float(k), epfd_db=float(-165 + 4 * rng.standard_normal()),
                duration_s=1.0, num_horizon_sats=8, num_visible_sats=4,
                num_contributing_sats=2, min_alpha_deg=4.0)
    return acc


def _sim_data() -> dict:
    bins = np.arange(-150.0, -180.1, -0.1)
    pct = np.geomspace(1e-4, 100.0, bins.size)
    return {
        "kind": "s1503",
        "epfd_type": "down",
        "input_source": "mdb",
        "compliance": "pass",
        "ccdf_bins_db": [round(float(b), 1) for b in bins],
        "ccdf_pct": [float(p) for p in pct],
        "article22": {
            "reference_bandwidth_khz": 40.0,
            "limits": [[-160.0, 0.1], [-175.0, 25.0]],
            "service": "FSS",
            "frequency_run_mhz": 17800.0,
        },
        "identification": {"ntc_id": "101", "sat_name": "TEST-SAT",
                            "mask_id": 3, "mask_source": "mdb"},
        "units": {"epfd": "dBW/m^2/40kHz"},
        "table17": [
            {"Ji_dBW": -160.0, "Pi_pct": 0.1, "Py_pct": 0.01, "pass": True},
            {"Ji_dBW": -175.0, "Pi_pct": 25.0, "Py_pct": 10.0, "pass": True},
        ],
    }


def _rows(path: Path) -> list[dict]:
    return list(csv.DictReader(l for l in path.read_text().splitlines()
                                if not l.startswith("#")))


def test_artifacts_written_and_consistent(tmp_path):
    sim = _sim_data()
    written = write_run_artifacts(tmp_path, sim, acc=_acc())
    assert {"ccdf_epfd.csv", "epfd_histogram.csv", "epfd_timeseries.csv",
            "table17.csv", "ccdf.png", "histogram.png"} <= set(written)

    ccdf = _rows(tmp_path / "ccdf_epfd.csv")
    assert len(ccdf) == len(sim["ccdf_bins_db"])
    assert float(ccdf[0]["epfd_db"]) == sim["ccdf_bins_db"][0]

    hist = _rows(tmp_path / "epfd_histogram.csv")
    assert abs(sum(float(r["probability"]) for r in hist) - 1.0) < 1e-9

    t17 = _rows(tmp_path / "table17.csv")
    assert [r["result"] for r in t17] == ["Pass", "Pass"]

    # units stated in headers (R23)
    for name in ("ccdf_epfd.csv", "epfd_histogram.csv", "table17.csv"):
        head = (tmp_path / name).read_text().splitlines()[1]
        assert "dBW/m^2/40kHz" in head


def test_table17_py_matches_ccdf():
    class FakeSim:
        cdf_epfd_dBW = np.arange(-150.0, -180.1, -0.1)
        cdf_percentage = np.linspace(0.001, 100.0, cdf_epfd_dBW.size)

    comp = check_article22_compliance(FakeSim, [(-165.0, 50.0)])
    row = comp.table17[0]
    # Py must equal the CCDF percentage at the last bin >= Ji.
    idx = int(np.searchsorted(-FakeSim.cdf_epfd_dBW, 165.0, side="right")) - 1
    assert row["Py_pct"] == pytest.approx(float(FakeSim.cdf_percentage[idx]))
    assert row["pass"] is (comp.worst_margin_dB >= 0.0)


def test_summary_html_blocks(tmp_path):
    sim = _sim_data()
    name = write_summary_html(tmp_path, sim)
    html = (tmp_path / name).read_text()
    for token in ("D7.3.1", "Table 17", "D7.3.3", "PASS", "TEST-SAT"):
        assert token in html


def test_run_to_xlsx_sheets():
    sim = _sim_data()
    sim["histogram"] = {"bin_size_db": 0.1, "bin_low_db": [-166.0],
                        "duration_s": [10.0], "probability": [1.0]}
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(run_to_xlsx(sim)))
    assert wb.sheetnames == ["run_def", "result_def", "results", "cdf", "pdf"]
    assert wb["cdf"].max_row == len(sim["ccdf_bins_db"]) + 1
    # run_def carries the identification (R18)
    vals = {r[0].value: r[1].value for r in wb["run_def"].iter_rows(min_row=2)}
    assert vals["sat_name"] == "TEST-SAT"
    assert vals["epfd_type"] == "down"


def test_geometries_from_per_point_and_map(tmp_path):
    from streamlit_app.lib.result_artifacts import (
        write_geometries_csv, write_geometry_map_png,
    )
    sim = _sim_data()
    sim["per_point"] = [
        {"index": 0, "es_lat_deg": -10.0, "es_lon_deg": -50.0,
         "gso_lon_deg": -45.0, "max_epfd_dbw": -160.0},
        {"index": 1, "es_lat_deg": 20.0, "es_lon_deg": 10.0,
         "gso_lon_deg": 5.0, "max_epfd_dbw": -158.5},
    ]
    assert write_geometries_csv(tmp_path, sim) == "geometries.csv"
    rows = _rows(tmp_path / "geometries.csv")
    assert len(rows) == 2 and rows[1]["measure_id"] == "1"
    assert write_geometry_map_png(tmp_path, sim) == "map.png"
    assert (tmp_path / "map.png").stat().st_size > 1000


def test_geometries_from_single_wcg(tmp_path):
    from streamlit_app.lib.result_artifacts import write_geometries_csv
    sim = _sim_data()
    sim["wcg"] = {"es_lat_deg": 5.45, "es_lon_deg": 0.16, "gso_lon_deg": 0.16}
    sim["max_epfd_dbw_m2_40khz"] = -155.9
    assert write_geometries_csv(tmp_path, sim) == "geometries.csv"
    rows = _rows(tmp_path / "geometries.csv")
    assert len(rows) == 1 and float(rows[0]["max_epfd_db"]) == -155.9


def test_timeseries_gz_for_long_traces(tmp_path):
    import gzip
    from streamlit_app.lib.result_artifacts import write_timeseries_csv

    class BigAcc:
        decim_t_s = list(range(60_000))
        decim_epfd_db = [-160.0] * 60_000
        decim_duration_s = [1.0] * 60_000
        decim_stride = 4
        n_steps = 240_000

    name = write_timeseries_csv(tmp_path, _sim_data(), BigAcc())
    assert name == "epfd_timeseries.csv.gz"
    text = gzip.decompress((tmp_path / name).read_bytes()).decode()
    assert text.count("\n") > 60_000


def test_run_to_xlsx_aggregate_curves_and_res76():
    """Aggregate export parity with the UI chart: per-point curves get their
    own sheet and Resolution 76 joins the result_def specification points."""
    import openpyxl
    sim = _sim_data()
    sim["method"] = "method_2"
    sim["resolution76"] = {"limits": [[-157.0, 0.1], [-170.0, 25.0]]}
    sim["per_point"] = [
        {"index": 0, "ccdf_bins_db": [-160.0, -170.0], "ccdf_pct": [0.01, 50.0]},
        {"index": 1, "ccdf_bins_db": [-158.0, -168.0, -175.0],
         "ccdf_pct": [0.005, 10.0, 90.0]},
    ]
    wb = openpyxl.load_workbook(io.BytesIO(run_to_xlsx(sim)))
    assert "cdf_per_point" in wb.sheetnames
    ws = wb["cdf_per_point"]
    heads = [c.value for c in ws[1]]
    assert "pt0_epfd [dBW/m^2/40kHz]" in heads and "pt1_pct [%]" in heads
    # ragged curves padded, all points preserved
    assert ws.max_row == 1 + 3
    rd = wb["result_def"]
    sets = {r[0].value for r in rd.iter_rows(min_row=2)}
    assert sets == {"Article 22", "Resolution 76 (aggregate)"}


def test_ccdf_png_includes_overlays(tmp_path):
    from streamlit_app.lib.result_artifacts import write_ccdf_png
    sim = _sim_data()
    sim["method"] = "method_2"
    sim["resolution76"] = {"limits": [[-157.0, 0.1], [-170.0, 25.0]]}
    sim["per_point"] = [
        {"index": 0, "ccdf_bins_db": [-160.0, -170.0], "ccdf_pct": [0.01, 50.0]},
    ]
    assert write_ccdf_png(tmp_path, sim) == "ccdf.png"
    assert (tmp_path / "ccdf.png").stat().st_size > 1000
